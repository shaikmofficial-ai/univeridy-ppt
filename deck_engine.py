#!/usr/bin/env python3
"""
Dual-output drawing engine for the UniVerify V2.0 academic project deck.

The same drawing code targets two backends:
  - PptxCanvas  -> editable PowerPoint (.pptx)  [the deliverable]
  - PdfCanvas   -> PDF (.pdf)                    [for visual QA + bonus copy]

Coordinate system: top-left origin, units in POINTS, page = 960 x 540 (16:9).
Text wrapping uses reportlab Helvetica metrics in BOTH backends so the line
breaks are identical (Calibri renders slightly narrower than Helvetica, so
pre-wrapped lines always fit inside the PPTX boxes).
"""

from reportlab.pdfbase.pdfmetrics import stringWidth

PAGE_W, PAGE_H = 960.0, 540.0
EMU_PER_PT = 12700


# ----------------------------------------------------------------- text utils
def _wrap(s, width, size, font="Helvetica"):
    """Wrap a string to a pixel width using Helvetica metrics."""
    out = []
    for para in s.split("\n"):
        words = para.split(" ")
        line = ""
        for w in words:
            trial = w if not line else line + " " + w
            if stringWidth(trial, font, size) <= width or not line:
                line = trial
            else:
                out.append(line)
                line = w
        out.append(line)
    return out


# ============================================================ PPTX BACKEND
class PptxCanvas:
    def __init__(self):
        from pptx import Presentation
        from pptx.util import Emu
        self._Emu = Emu
        self.prs = Presentation()
        self.prs.slide_width = Emu(int(PAGE_W * EMU_PER_PT))
        self.prs.slide_height = Emu(int(PAGE_H * EMU_PER_PT))
        self._blank = self.prs.slide_layouts[6]
        self.slide = None

    # --- helpers ---
    def _E(self, v):
        return self._Emu(int(round(v * EMU_PER_PT)))

    def _rgb(self, hexs):
        from pptx.dml.color import RGBColor
        return RGBColor.from_string(hexs.lstrip("#"))

    def _style(self, shp, fill, line, line_w):
        shp.shadow.inherit = False
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = self._rgb(fill)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = self._rgb(line)
            from pptx.util import Pt
            shp.line.width = Pt(line_w)

    # --- API ---
    def new_page(self, bg="#FFFFFF"):
        self.slide = self.prs.slides.add_slide(self._blank)
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = self._rgb(bg)

    def rect(self, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0):
        from pptx.enum.shapes import MSO_SHAPE
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shp = self.slide.shapes.add_shape(kind, self._E(x), self._E(y),
                                          self._E(w), self._E(h))
        if radius:
            try:
                shp.adjustments[0] = max(0.0, min(0.5, radius / min(w, h)))
            except Exception:
                pass
        self._style(shp, fill, line, line_w)
        return shp

    def ellipse(self, x, y, w, h, fill=None, line=None, line_w=1.0):
        from pptx.enum.shapes import MSO_SHAPE
        shp = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, self._E(x), self._E(y),
                                          self._E(w), self._E(h))
        self._style(shp, fill, line, line_w)
        return shp

    def diamond(self, x, y, w, h, fill=None, line=None, line_w=1.0):
        from pptx.enum.shapes import MSO_SHAPE
        shp = self.slide.shapes.add_shape(MSO_SHAPE.DIAMOND, self._E(x), self._E(y),
                                          self._E(w), self._E(h))
        self._style(shp, fill, line, line_w)
        return shp

    def line(self, x1, y1, x2, y2, color="#000000", w=1.0,
             arrow_end=False, arrow_start=False, dash=False):
        from pptx.enum.shapes import MSO_CONNECTOR
        from pptx.util import Pt
        from pptx.oxml.ns import qn
        conn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                               self._E(x1), self._E(y1),
                                               self._E(x2), self._E(y2))
        conn.line.color.rgb = self._rgb(color)
        conn.line.width = Pt(w)
        try:
            conn.shadow.inherit = False
        except Exception:
            pass
        spPr = conn._element.find(qn("p:spPr"))
        ln = spPr.find(qn("a:ln"))
        if ln is not None:
            if dash:
                d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
                ln.append(d)
            if arrow_start:
                he = ln.makeelement(qn("a:headEnd"),
                                    {"type": "triangle", "w": "med", "len": "med"})
                ln.append(he)
            if arrow_end:
                te = ln.makeelement(qn("a:tailEnd"),
                                    {"type": "triangle", "w": "med", "len": "med"})
                ln.append(te)
        return conn

    def text(self, x, y, w, h, s, size=12, color="#16213E", align="l",
             valign="t", bold=False, italic=False, font="Calibri", wrap=True):
        from pptx.util import Pt
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        tb = self.slide.shapes.add_textbox(self._E(x), self._E(y),
                                           self._E(w), self._E(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                              "b": MSO_ANCHOR.BOTTOM}[valign]
        al = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        for i, ln in enumerate(s.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = al
            p.line_spacing = 1.0
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            r = p.add_run()
            r.text = ln
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.name = font
            r.font.color.rgb = self._rgb(color)
        return tb

    def save(self, path):
        self.prs.save(path)


# ============================================================ PDF BACKEND
class PdfCanvas:
    def __init__(self, path):
        from reportlab.pdfgen import canvas
        self.path = path
        self.c = canvas.Canvas(path, pagesize=(PAGE_W, PAGE_H))
        self._open = False

    def _col(self, hexs):
        from reportlab.lib.colors import HexColor
        return HexColor(hexs)

    def new_page(self, bg="#FFFFFF"):
        if self._open:
            self.c.showPage()
        self._open = True
        self.c.setFillColor(self._col(bg))
        self.c.rect(0, 0, PAGE_W, PAGE_H, stroke=0, fill=1)

    def _draw(self, fill, line, line_w):
        do_fill = 0
        do_stroke = 0
        if fill is not None:
            self.c.setFillColor(self._col(fill))
            do_fill = 1
        if line is not None:
            self.c.setStrokeColor(self._col(line))
            self.c.setLineWidth(line_w)
            do_stroke = 1
        return do_fill, do_stroke

    def rect(self, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0):
        do_fill, do_stroke = self._draw(fill, line, line_w)
        ry = PAGE_H - y - h
        if radius:
            self.c.roundRect(x, ry, w, h, radius, stroke=do_stroke, fill=do_fill)
        else:
            self.c.rect(x, ry, w, h, stroke=do_stroke, fill=do_fill)

    def ellipse(self, x, y, w, h, fill=None, line=None, line_w=1.0):
        do_fill, do_stroke = self._draw(fill, line, line_w)
        ry = PAGE_H - y - h
        self.c.ellipse(x, ry, x + w, ry + h, stroke=do_stroke, fill=do_fill)

    def diamond(self, x, y, w, h, fill=None, line=None, line_w=1.0):
        do_fill, do_stroke = self._draw(fill, line, line_w)
        cx, cy = x + w / 2.0, PAGE_H - (y + h / 2.0)
        p = self.c.beginPath()
        p.moveTo(cx, cy + h / 2.0)
        p.lineTo(x + w, cy)
        p.lineTo(cx, cy - h / 2.0)
        p.lineTo(x, cy)
        p.close()
        self.c.drawPath(p, stroke=do_stroke, fill=do_fill)

    def line(self, x1, y1, x2, y2, color="#000000", w=1.0,
             arrow_end=False, arrow_start=False, dash=False):
        import math
        self.c.setStrokeColor(self._col(color))
        self.c.setFillColor(self._col(color))
        self.c.setLineWidth(w)
        if dash:
            self.c.setDash(4, 3)
        Y1, Y2 = PAGE_H - y1, PAGE_H - y2
        self.c.line(x1, Y1, x2, Y2)
        if dash:
            self.c.setDash([])

        def head(xa, ya, xb, yb):
            ang = math.atan2(yb - ya, xb - xa)
            sz = 7
            a1 = ang + math.radians(150)
            a2 = ang - math.radians(150)
            p = self.c.beginPath()
            p.moveTo(xb, yb)
            p.lineTo(xb + sz * math.cos(a1), yb + sz * math.sin(a1))
            p.lineTo(xb + sz * math.cos(a2), yb + sz * math.sin(a2))
            p.close()
            self.c.drawPath(p, stroke=0, fill=1)

        if arrow_end:
            head(x1, Y1, x2, Y2)
        if arrow_start:
            head(x2, Y2, x1, Y1)

    def text(self, x, y, w, h, s, size=12, color="#16213E", align="l",
             valign="t", bold=False, italic=False, font="Helvetica", wrap=True):
        fname = "Helvetica"
        if bold and italic:
            fname = "Helvetica-BoldOblique"
        elif bold:
            fname = "Helvetica-Bold"
        elif italic:
            fname = "Helvetica-Oblique"
        self.c.setFillColor(self._col(color))
        self.c.setFont(fname, size)
        lines = s.split("\n")
        leading = size * 1.25
        total = len(lines) * leading
        if valign == "m":
            top = y + (h - total) / 2.0
        elif valign == "b":
            top = y + (h - total)
        else:
            top = y
        for i, ln in enumerate(lines):
            baseline = PAGE_H - (top + i * leading + size * 0.82)
            if align == "c":
                self.c.drawCentredString(x + w / 2.0, baseline, ln)
            elif align == "r":
                self.c.drawRightString(x + w, baseline, ln)
            else:
                self.c.drawString(x, baseline, ln)

    def save(self, path=None):
        if self._open:
            self.c.showPage()
        self.c.save()
